"""Node B: ingestion into ChromaDB.

Reads the reports in `data/synth_reports/` (Markdown/Word/Excel/PowerPoint/PDF), chunks them by
section, embeds them with a local model, and stores them in ChromaDB.

- Metadata comes from the structured header datagen writes (Markdown front matter, the first Word
  table, the first Excel rows, the first slide, or the leading PDF text), never from free text.
- Chunks follow `## heading` sections, not character counts. Word/PowerPoint structure is reliable;
  PDF has none, so it relies on text patterns, as real-world PDF ingestion does.
- The one result image in each report's result section is captioned by the VLM
  (config.ai.vlm_model) and appended to that chunk so it is searchable. A captioning failure only
  drops that caption; ingestion continues.
"""

from __future__ import annotations

import argparse
import contextlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .config import CHROMA_DIR, COLLECTION_NAME, SYNTH_REPORTS_DIR, load_config
from .llm_client import LLMClient, LLMConnectionError

_FRONTMATTER_RE = re.compile(r"\A---\n(.*?)\n---\n", re.DOTALL)
_SECTION_RE = re.compile(r"^## +(.+?)\s*$", re.MULTILINE)
_MARKDOWN_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")

# Same value as datagen.RESULT_IMAGE_SECTION, duplicated to avoid depending on datagen
# (against the DAG direction), like the heading regex above.
RESULT_IMAGE_SECTION = "成果サマリー"

_IMAGE_CAPTION_PROMPT = (
    "これは社内のプロジェクト振り返りレポートに含まれる成果画像（グラフ）です。"
    "何を表しているか、特徴的な傾向（ピーク位置・値の大小・推移の変化など）を"
    "2〜3文の日本語で具体的に説明してください。"
)


@dataclass
class Chunk:
    chunk_id: str
    text: str
    metadata: dict[str, str]


def parse_frontmatter(raw: str) -> tuple[dict[str, str], str]:
    """Parse the leading `---` front matter; return (metadata, remaining body)."""
    m = _FRONTMATTER_RE.match(raw)
    if not m:
        raise ValueError(
            "フロントマター（--- ... ---）が見つかりません。datagenで生成したファイルか確認してください。"
        )
    meta: dict[str, str] = {}
    for line in m.group(1).splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip()
    body = raw[m.end() :]
    return meta, body


def split_into_sections(body: str) -> list[tuple[str, str]]:
    """Split a Markdown body at each `## heading`, ignoring the title before the first one."""
    matches = list(_SECTION_RE.finditer(body))
    sections: list[tuple[str, str]] = []
    for i, m in enumerate(matches):
        heading = m.group(1).strip()
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
        text = body[start:end].strip()
        if text:
            sections.append((heading, text))
    return sections


def _extract_markdown(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    raw = path.read_text(encoding="utf-8")
    meta, body = parse_frontmatter(raw)
    sections = split_into_sections(body)

    # datagen._write_markdown embeds the image as `![alt](file)`: strip the link, read the sibling PNG.
    image: bytes | None = None
    cleaned: list[tuple[str, str]] = []
    for heading, text in sections:
        m = _MARKDOWN_IMAGE_RE.search(text)
        if m:
            image_path = path.parent / m.group(1)
            if image_path.is_file():
                image = image_path.read_bytes()
            text = _MARKDOWN_IMAGE_RE.sub("", text).strip()
        cleaned.append((heading, text))
    return meta, cleaned, image


def _extract_docx(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    """Read metadata from the first table and sections from heading styles in a Word file."""
    from docx import Document

    doc = Document(str(path))
    if not doc.tables:
        raise ValueError(
            f"{path}: メタデータの表が見つかりません。datagenで生成したファイルか確認してください。"
        )
    meta = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in doc.tables[0].rows}

    sections: list[tuple[str, str]] = []
    current_heading: str | None = None
    current_lines: list[str] = []
    for para in doc.paragraphs:
        style_name = para.style.name if para.style is not None else ""
        # "Heading 1" is the title; "Heading 2" and below are sections.
        if style_name.startswith("Heading") and style_name != "Heading 1":
            if current_heading is not None:
                sections.append((current_heading, "\n".join(current_lines).strip()))
            current_heading = para.text.strip()
            current_lines = []
        elif current_heading is not None:
            current_lines.append(para.text)
    if current_heading is not None:
        sections.append((current_heading, "\n".join(current_lines).strip()))

    # A report has at most one image, so take the first image relationship. InlineShape has no
    # `.image` attribute in python-docx, so go through the document relationships instead.
    image: bytes | None = None
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                image = rel.target_part.blob
                break
    except Exception as exc:
        print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
    return meta, sections, image


def _extract_xlsx(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    """Read metadata and (heading, text) section rows from an Excel file.

    Metadata runs until the first empty cell in column A; the next row is a separator.
    """
    from openpyxl import load_workbook

    # read_only mode doesn't load ws._images; memory isn't a concern at this data size.
    wb = load_workbook(str(path), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    meta: dict[str, str] = {}
    idx = 0
    while idx < len(rows) and rows[idx] and rows[idx][0]:
        key = str(rows[idx][0])
        value = rows[idx][1] if len(rows[idx]) > 1 else None
        meta[key] = "" if value is None else str(value)
        idx += 1
    idx += 1  # skip the blank separator row

    sections: list[tuple[str, str]] = []
    for row in rows[idx:]:
        if not row or not row[0]:
            continue
        heading = str(row[0])
        value = row[1] if len(row) > 1 else None
        sections.append((heading, "" if value is None else str(value)))

    # `.ref` turned out to be a BytesIO, not a PIL Image, so use `_data()` (the raw bytes openpyxl
    # itself writes) instead of relying on `.ref`'s internal type.
    image: bytes | None = None
    embedded_images = getattr(ws, "_images", [])
    if embedded_images:
        try:
            image = embedded_images[0]._data()
        except Exception as exc:
            print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
    return meta, sections, image


def _extract_pptx(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    """Read a PowerPoint file: slide 1 is metadata, each later slide is a section."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    slides = list(prs.slides)
    if not slides:
        raise ValueError(f"{path}: スライドが1枚もありません。")

    def _body_text(slide) -> str:
        for shape in slide.placeholders:
            # idx 0 is the title placeholder; anything else is body.
            if shape.placeholder_format.idx != 0 and shape.has_text_frame:
                return shape.text_frame.text
        return ""

    meta: dict[str, str] = {}
    for line in _body_text(slides[0]).splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip()

    sections: list[tuple[str, str]] = []
    image: bytes | None = None
    for slide in slides[1:]:
        heading = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        sections.append((heading, _body_text(slide)))
        if image is None:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image.blob
                    except Exception as exc:
                        print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
                    break
    return meta, sections, image


def _parse_pdf_text(full_text: str) -> tuple[dict[str, str], list[tuple[str, str]]]:
    """Parse extracted PDF text: `key: value` lines are metadata, `## heading` lines start sections.

    Assumes the layout from datagen._pdf_lines(). A pure function (no pypdf), so the round trip
    can be tested without pypdf or reportlab.

    Metadata ends at the first `## ` line, not at a blank line: pypdf can drop blank lines that were
    only a y-offset, which once made every section unreadable (9 of 60 PDFs).
    """
    lines = full_text.split("\n")

    meta: dict[str, str] = {}
    sections: list[tuple[str, str]] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("## "):
            if current_heading is not None:
                sections.append((current_heading, "\n".join(current_lines).strip()))
            current_heading = stripped[3:].strip()
            current_lines = []
        elif current_heading is None:
            # Before the first heading: metadata.
            if ":" in stripped:
                key, _, value = stripped.partition(":")
                meta[key.strip()] = value.strip()
        elif stripped:
            current_lines.append(stripped)

    if current_heading is not None:
        sections.append((current_heading, "\n".join(current_lines).strip()))
    return meta, sections


def _extract_pdf(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    meta, sections = _parse_pdf_text(full_text)

    image: bytes | None = None
    for page in reader.pages:
        page_images: Any
        try:
            page_images = page.images
        except Exception:
            page_images = []
        if page_images:
            try:
                image = page_images[0].data
            except Exception as exc:
                print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
            break
    return meta, sections, image


_EXTRACTORS = {
    ".md": _extract_markdown,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
    ".pdf": _extract_pdf,
}


def _caption_image(client: LLMClient, image: bytes) -> str | None:
    """Caption an image with the VLM, or return None on failure (e.g. VLM not loaded).

    Like retrieval.py's rerank fallback: a failure only drops this caption, ingestion continues.
    """
    try:
        return client.describe_image(image, _IMAGE_CAPTION_PROMPT)
    except LLMConnectionError as exc:
        print(f"画像キャプション取得に失敗しました（無視して続行）: {exc}")
        return None


def _find_report_paths(reports_dir: Path) -> list[Path]:
    return sorted(p for ext in _EXTRACTORS for p in reports_dir.glob(f"*{ext}"))


def build_chunks(report_path: Path, client: LLMClient | None = None) -> list[Chunk]:
    """Build the chunks for one file.

    With a client, an embedded image is captioned and appended to the result-section chunk so it is
    searchable. client=None skips captioning (e.g. tests without an LLM).
    """
    extractor = _EXTRACTORS.get(report_path.suffix.lower())
    if extractor is None:
        raise ValueError(f"未対応のファイル形式です: {report_path}")
    meta, sections, image = extractor(report_path)
    report_id = meta.get("report_id", report_path.stem)

    caption: str | None = None
    if image is not None and client is not None:
        caption = _caption_image(client, image)

    chunks: list[Chunk] = []
    for heading, text in sections:
        if heading == RESULT_IMAGE_SECTION and caption:
            text = f"{text}\n\n[結果画像の説明] {caption}" if text else f"[結果画像の説明] {caption}"
        if not text:
            continue
        chunk_id = f"{report_id}::{heading}"
        chunk_text = f"【{heading}】\n{text}"
        chunk_meta = {**meta, "section": heading}
        chunks.append(Chunk(chunk_id=chunk_id, text=chunk_text, metadata=chunk_meta))
    return chunks


def load_all_chunks(reports_dir: Path = SYNTH_REPORTS_DIR, client: LLMClient | None = None) -> list[Chunk]:
    chunks: list[Chunk] = []
    for path in _find_report_paths(reports_dir):
        chunks.extend(build_chunks(path, client))
    return chunks


def ingest(reports_dir: Path = SYNTH_REPORTS_DIR, chroma_dir: Path = CHROMA_DIR) -> int:
    """Embed the reports into ChromaDB and return the number of chunks stored."""
    import chromadb

    report_paths = _find_report_paths(reports_dir)
    if not report_paths:
        hint = "先にdatagenを実行してください。" if reports_dir == SYNTH_REPORTS_DIR else ""
        raise SystemExit(f"{reports_dir} にレポートが見つかりません。{hint}")

    config = load_config()
    chroma_dir.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(chroma_dir))

    with LLMClient(config.ai) as llm_client:
        if not llm_client.ping():
            raise SystemExit(
                f"LM Studio ({config.ai.base_url}) に接続できません。起動してモデルをロードしてください。"
            )

        # Building chunks also captions images, using the same client as embedding.
        chunks: list[Chunk] = []
        for path in report_paths:
            chunks.extend(build_chunks(path, llm_client))
        if not chunks:
            raise SystemExit(f"{reports_dir} のレポートからチャンクを構築できませんでした。")

        # Build into a temporary collection and swap only after every embedding succeeds, so a
        # failure never breaks the existing index. list_collections() returns different types across
        # chromadb versions, so just try deleting instead of checking existence.
        building_name = f"{COLLECTION_NAME}__building"
        _delete_collection_if_exists(client, building_name)
        building = client.create_collection(building_name)

        try:
            # Batch so each embedding request stays a reasonable size.
            batch_size = 32
            for i in range(0, len(chunks), batch_size):
                batch = chunks[i : i + batch_size]
                embeddings = llm_client.embed([c.text for c in batch])
                building.add(
                    ids=[c.chunk_id for c in batch],
                    embeddings=embeddings,  # type: ignore[arg-type]  # list invariance vs. chromadb stub
                    documents=[c.text for c in batch],
                    metadatas=[c.metadata for c in batch],
                )
                print(f"embedded {min(i + batch_size, len(chunks))}/{len(chunks)} chunks")
        except Exception:
            _delete_collection_if_exists(client, building_name)
            raise

        _delete_collection_if_exists(client, COLLECTION_NAME)
        building.modify(name=COLLECTION_NAME)

    return len(chunks)


def _delete_collection_if_exists(client, name: str) -> None:
    with contextlib.suppress(Exception):
        client.delete_collection(name)


def main() -> None:
    parser = argparse.ArgumentParser(description="レポートをチャンキング・埋め込みしてChromaDBに格納")
    parser.add_argument(
        "--reports-dir",
        type=Path,
        default=SYNTH_REPORTS_DIR,
        help="読み込むレポートのディレクトリ（既定: datagenの出力先）。自前データを使う場合はここを指定する",
    )
    args = parser.parse_args()

    count = ingest(reports_dir=args.reports_dir)
    print(f"ingest完了: {count} チャンクを {CHROMA_DIR} に格納しました。")


if __name__ == "__main__":
    main()
