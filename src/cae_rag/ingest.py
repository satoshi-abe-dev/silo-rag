"""Ingestion パイプライン（DAGノードB）。

`data/synth_reports/*`（datagenノードが生成したダミー解析レポート。Markdown/Word/Excel/
PowerPoint/PDFのいずれか）を読み込み、セクション単位でチャンキングし、ローカル埋め込み
モデルでベクトル化してChromaDBに格納する。

メタデータ（部署・解析種別・部品・ソルバー・日付等）は、datagenが各ファイルの先頭に書き出す
ヘッダーブロック（Markdownなら`---`フロントマター、Wordなら先頭の表、Excelなら先頭の行、
PowerPointなら1枚目のスライド、PDFなら先頭のテキスト）から直接読み取る。自由記述文からの
あいまいな抽出は行わない（レポート形式を自分たちで決められるため、構造化ヘッダーの方が
確実で、実務の社内文書でもよくあるパターンでもある）。

チャンキングは `## 見出し` 単位（セクション境界を尊重し、単純な文字数分割はしない）。
ファイル形式によって、その構造をどこまで正確に読み取れるかには差がある（Word/PowerPointは
見出しスタイルやスライド構造から確実に判定できるが、PDFは本質的に構造を持たないため
テキストパターン頼みの抽出になる。実務のPDF取り込みでよくある制約をそのまま反映している）。

各レポートには「結果サマリー」セクションに1枚だけ結果画像（グラフ/コンター図）が
埋め込まれている想定。ファイル形式ごとの方法で画像を取り出し、LM StudioのVLM
（config.llm.vlm_model）でキャプション化して、対応するチャンクの検索対象テキストに
追記する（画像の内容も検索できるようにするため）。VLM未ロード等で失敗しても、
その画像のキャプションが無いだけでテキスト取り込み自体は続行する。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from .config import CHROMA_DIR, COLLECTION_NAME, SYNTH_REPORTS_DIR, load_config
from .llm_client import LLMClient, LLMConnectionError

_FRONTMATTER_RE = re.compile(r"\A---\n(.*?)\n---\n", re.DOTALL)
_SECTION_RE = re.compile(r"^## +(.+?)\s*$", re.MULTILINE)
_MARKDOWN_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")

# 画像が必ず添付されるセクション。datagen.RESULT_IMAGE_SECTIONと同じ値を独自に持つ
# （datagenへ依存する逆向きの結合を避けるため。既存の見出し正規表現と同じ方針）。
RESULT_IMAGE_SECTION = "結果サマリー"

_IMAGE_CAPTION_PROMPT = (
    "これは社内のCAE構造解析レポートに含まれる結果画像（コンター図またはグラフ）です。"
    "何を表しているか、特徴的な傾向（ピーク位置・値の大小・分布の偏りなど）を"
    "2〜3文の日本語で具体的に説明してください。"
)


@dataclass
class Chunk:
    chunk_id: str
    text: str
    metadata: dict[str, str]


def parse_frontmatter(raw: str) -> tuple[dict[str, str], str]:
    """先頭の `---` フロントマターを辞書として取り出し、残りの本文と一緒に返す。"""
    m = _FRONTMATTER_RE.match(raw)
    if not m:
        raise ValueError("フロントマター（--- ... ---）が見つかりません。datagenで生成したファイルか確認してください。")
    meta: dict[str, str] = {}
    for line in m.group(1).splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip()
    body = raw[m.end() :]
    return meta, body


def split_into_sections(body: str) -> list[tuple[str, str]]:
    """`## 見出し` ごとにMarkdown本文を分割する。見出し前のタイトル行は無視する。"""
    matches = list(_SECTION_RE.finditer(body))
    sections: list[tuple[str, str]] = []
    for i, m in enumerate(matches):
        heading = m.group(1).strip()
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
        text = body[start:end].strip()
        if text:
            sections.append((heading, text))
    return sections


def _extract_markdown(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    raw = path.read_text(encoding="utf-8")
    meta, body = parse_frontmatter(raw)
    sections = split_into_sections(body)

    # 画像は `![alt](ファイル名)` としてセクション本文中に埋め込まれている
    # （datagen._write_markdown参照）。見つけたら本文から取り除き、隣のPNGファイルを読む。
    image: bytes | None = None
    cleaned: list[tuple[str, str]] = []
    for heading, text in sections:
        m = _MARKDOWN_IMAGE_RE.search(text)
        if m:
            image_path = path.parent / m.group(1)
            if image_path.is_file():
                image = image_path.read_bytes()
            text = _MARKDOWN_IMAGE_RE.sub("", text).strip()
        cleaned.append((heading, text))
    return meta, cleaned, image


def _extract_docx(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    """Wordファイルからメタデータ（先頭の表）とセクション（見出しスタイル単位）を取り出す。"""
    from docx import Document

    doc = Document(str(path))
    if not doc.tables:
        raise ValueError(f"{path}: メタデータの表が見つかりません。datagenで生成したファイルか確認してください。")
    meta = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in doc.tables[0].rows}

    sections: list[tuple[str, str]] = []
    current_heading: str | None = None
    current_lines: list[str] = []
    for para in doc.paragraphs:
        style_name = para.style.name if para.style is not None else ""
        # "Heading 1" はタイトル行なので対象外。"Heading 2" 以降がセクション見出し。
        if style_name.startswith("Heading") and style_name != "Heading 1":
            if current_heading is not None:
                sections.append((current_heading, "\n".join(current_lines).strip()))
            current_heading = para.text.strip()
            current_lines = []
        elif current_heading is not None:
            current_lines.append(para.text)
    if current_heading is not None:
        sections.append((current_heading, "\n".join(current_lines).strip()))

    # レポート1件につき画像は最大1枚（datagen._write_docx参照）なので、
    # 最初に見つかった画像リレーションシップをそのまま結果画像として扱う。
    # InlineShapeには`.image`属性が無い（python-docxの実際のAPIと異なっていた）ため、
    # ドキュメントのリレーションシップから直接画像パートを辿る。
    image: bytes | None = None
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                image = rel.target_part.blob
                break
    except Exception as exc:  # noqa: BLE001 - 画像取得の失敗はテキスト取り込みを止めない
        print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
    return meta, sections, image


def _extract_xlsx(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    """Excelファイルからメタデータ（先頭の行群）とセクション（見出し, 内容の行）を取り出す。

    行数を決め打ちにせず、A列が空になる行までをメタデータとみなし、
    その次の1行を区切りとしてスキップしてからセクション行を読む。
    """
    from openpyxl import load_workbook

    # 画像（ws._images）はread_onlyモードでは読み込まれないため、通常モードで開く
    # （このプロジェクトのデータ規模ではメモリ効率を気にする必要はない）。
    wb = load_workbook(str(path), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    meta: dict[str, str] = {}
    idx = 0
    while idx < len(rows) and rows[idx] and rows[idx][0]:
        key = str(rows[idx][0])
        value = rows[idx][1] if len(rows[idx]) > 1 else None
        meta[key] = "" if value is None else str(value)
        idx += 1
    idx += 1  # 区切りの空行をスキップ

    sections: list[tuple[str, str]] = []
    for row in rows[idx:]:
        if not row or not row[0]:
            continue
        heading = str(row[0])
        value = row[1] if len(row) > 1 else None
        sections.append((heading, "" if value is None else str(value)))

    # `.ref`はPIL Imageだと想定していたが、実際にはBytesIO（`.save()`を持たない）で
    # 返ってくることが確認された。`_data()`はopenpyxl自身が書き出し時に使う、生バイト列を
    # 返すメソッドなので、`.ref`の内部表現に依存せずこちらを使う。
    image: bytes | None = None
    embedded_images = getattr(ws, "_images", [])
    if embedded_images:
        try:
            image = embedded_images[0]._data()
        except Exception as exc:  # noqa: BLE001 - 画像取得の失敗はテキスト取り込みを止めない
            print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
    return meta, sections, image


def _extract_pptx(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    """PowerPointファイルから、1枚目=メタデータ、2枚目以降=セクションとして取り出す。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    slides = list(prs.slides)
    if not slides:
        raise ValueError(f"{path}: スライドが1枚もありません。")

    def _body_text(slide) -> str:
        for shape in slide.placeholders:
            # placeholder_format.idx == 0 はタイトル。それ以外を本文とみなす。
            if shape.placeholder_format.idx != 0 and shape.has_text_frame:
                return shape.text_frame.text
        return ""

    meta: dict[str, str] = {}
    for line in _body_text(slides[0]).splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip()

    sections: list[tuple[str, str]] = []
    image: bytes | None = None
    for slide in slides[1:]:
        heading = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        sections.append((heading, _body_text(slide)))
        if image is None:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image.blob
                    except Exception as exc:  # noqa: BLE001
                        print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
                    break
    return meta, sections, image


def _parse_pdf_text(full_text: str) -> tuple[dict[str, str], list[tuple[str, str]]]:
    """PDFから抽出済みのプレーンテキストを、`key: value` 行をメタデータ、`## 見出し` 行を
    セクション境界としてパースする（pypdfに依存しない純粋関数）。

    datagen._pdf_lines() が組み立てる行の並びを前提にしている。PDFは本質的に構造を
    持たないため、他形式と違いテキストパターン頼みの抽出になる（実務のPDF取り込みでも
    よくある制約）。pypdfからのテキスト抽出処理と分離してあるので、pypdf/reportlabが
    無い環境でも datagen._pdf_lines() の出力を直接渡して往復ロジックをテストできる。
    """
    lines = full_text.split("\n")

    meta: dict[str, str] = {}
    idx = 0
    while idx < len(lines) and lines[idx].strip():
        line = lines[idx].strip()
        if ":" in line:
            key, _, value = line.partition(":")
            meta[key.strip()] = value.strip()
        idx += 1
    idx += 1  # 区切りの空行をスキップ

    sections: list[tuple[str, str]] = []
    current_heading: str | None = None
    current_lines: list[str] = []
    for line in lines[idx:]:
        stripped = line.strip()
        if stripped.startswith("## "):
            if current_heading is not None:
                sections.append((current_heading, "\n".join(current_lines).strip()))
            current_heading = stripped[3:].strip()
            current_lines = []
        elif current_heading is not None and stripped:
            current_lines.append(stripped)
    if current_heading is not None:
        sections.append((current_heading, "\n".join(current_lines).strip()))
    return meta, sections


def _extract_pdf(path: Path) -> tuple[dict[str, str], list[tuple[str, str]], bytes | None]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    meta, sections = _parse_pdf_text(full_text)

    image: bytes | None = None
    for page in reader.pages:
        try:
            page_images = page.images
        except Exception:  # noqa: BLE001
            page_images = []
        if page_images:
            try:
                image = page_images[0].data
            except Exception as exc:  # noqa: BLE001
                print(f"{path}: 画像の取得に失敗しました（無視して続行）: {exc}")
            break
    return meta, sections, image


_EXTRACTORS = {
    ".md": _extract_markdown,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
    ".pdf": _extract_pdf,
}


def _caption_image(client: LLMClient, image: bytes) -> str | None:
    """画像をVLMでキャプション化する。

    VLM未ロード等で失敗しても、テキスト取り込み全体は止めない
    （retrieval.pyのLLMリランク失敗時のフォールバックと同じ考え方。この画像1枚の
    キャプションが無いままチャンクを作るだけで、レポート全体の取り込みは続行する）。
    """
    try:
        return client.describe_image(image, _IMAGE_CAPTION_PROMPT)
    except LLMConnectionError as exc:
        print(f"画像キャプション取得に失敗しました（無視して続行）: {exc}")
        return None


def _find_report_paths(reports_dir: Path) -> list[Path]:
    return sorted(p for ext in _EXTRACTORS for p in reports_dir.glob(f"*{ext}"))


def build_chunks(report_path: Path, client: LLMClient | None = None) -> list[Chunk]:
    """1ファイル分のチャンクを作る。

    client を渡すと、埋め込み画像があればVLMでキャプション化し、結果セクションの
    チャンク本文に追記する（画像の内容も検索対象にするため）。client=Noneなら
    キャプション化をスキップする（テスト等、LLM接続なしでの利用を想定）。
    """
    extractor = _EXTRACTORS.get(report_path.suffix.lower())
    if extractor is None:
        raise ValueError(f"未対応のファイル形式です: {report_path}")
    meta, sections, image = extractor(report_path)
    report_id = meta.get("report_id", report_path.stem)

    caption: str | None = None
    if image is not None and client is not None:
        caption = _caption_image(client, image)

    chunks: list[Chunk] = []
    for heading, text in sections:
        if heading == RESULT_IMAGE_SECTION and caption:
            text = f"{text}\n\n[結果画像の説明] {caption}" if text else f"[結果画像の説明] {caption}"
        if not text:
            continue
        chunk_id = f"{report_id}::{heading}"
        chunk_text = f"【{heading}】\n{text}"
        chunk_meta = {**meta, "section": heading}
        chunks.append(Chunk(chunk_id=chunk_id, text=chunk_text, metadata=chunk_meta))
    return chunks


def load_all_chunks(reports_dir: Path = SYNTH_REPORTS_DIR, client: LLMClient | None = None) -> list[Chunk]:
    chunks: list[Chunk] = []
    for path in _find_report_paths(reports_dir):
        chunks.extend(build_chunks(path, client))
    return chunks


def ingest(reports_dir: Path = SYNTH_REPORTS_DIR, chroma_dir: Path = CHROMA_DIR) -> int:
    """レポートを読み込み、埋め込みを計算してChromaDBに格納する。格納したチャンク数を返す。"""
    import chromadb

    report_paths = _find_report_paths(reports_dir)
    if not report_paths:
        raise SystemExit(f"{reports_dir} にレポートが見つかりません。先にdatagenを実行してください。")

    config = load_config()
    chroma_dir.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(chroma_dir))

    with LLMClient(config.llm) as llm_client:
        if not llm_client.ping():
            raise SystemExit(
                f"LM Studio ({config.llm.base_url}) に接続できません。起動してモデルをロードしてください。"
            )

        # チャンク構築（埋め込み画像があればここでVLMキャプション化も行う）は、
        # 埋め込み・格納に使うのと同じLLMClientインスタンスで行う。
        chunks: list[Chunk] = []
        for path in report_paths:
            chunks.extend(build_chunks(path, llm_client))
        if not chunks:
            raise SystemExit(f"{reports_dir} のレポートからチャンクを構築できませんでした。")

        # 新しいコレクションを一時名で組み立て、全チャンクの埋め込みが成功してから
        # 旧コレクションと入れ替える。途中で失敗しても既存の検索用インデックスを壊さない。
        # list_collections()の返り値の形（Collectionオブジェクト/名前の文字列）は
        # chromadbのバージョンによって異なるため、存在チェックはせず削除を試みて無視する。
        building_name = f"{COLLECTION_NAME}__building"
        _delete_collection_if_exists(client, building_name)
        building = client.create_collection(building_name)

        try:
            # 大量チャンクでも1リクエストに収まるよう、適度なバッチサイズで埋め込みを取得する。
            batch_size = 32
            for i in range(0, len(chunks), batch_size):
                batch = chunks[i : i + batch_size]
                embeddings = llm_client.embed([c.text for c in batch])
                building.add(
                    ids=[c.chunk_id for c in batch],
                    embeddings=embeddings,
                    documents=[c.text for c in batch],
                    metadatas=[c.metadata for c in batch],
                )
                print(f"embedded {min(i + batch_size, len(chunks))}/{len(chunks)} chunks")
        except Exception:
            _delete_collection_if_exists(client, building_name)
            raise

        _delete_collection_if_exists(client, COLLECTION_NAME)
        building.modify(name=COLLECTION_NAME)

    return len(chunks)


def _delete_collection_if_exists(client, name: str) -> None:
    try:
        client.delete_collection(name)
    except Exception:
        pass


def main() -> None:
    count = ingest()
    print(f"ingest完了: {count} チャンクを {CHROMA_DIR} に格納しました。")


if __name__ == "__main__":
    main()
